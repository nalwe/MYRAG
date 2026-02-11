from django.db import transaction
from django.urls import reverse
from django.http import HttpResponse, HttpResponseForbidden
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.db.models import Q

from reportlab.pdfgen import canvas
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet
from io import BytesIO
from bs4 import BeautifulSoup
from docx import Document

import markdown

from rag.models import ChatSession, ChatMessage
from documents.models import Document as UserDocument
from documents.utils import get_accessible_documents
from rag.utils import create_onboarding_chat
from rag.retrieval import retrieve_chunks_for_chat
from rag.rag_pipeline import rag_answer

from accounts.models import OrganizationMember


# =====================================================
# 🔐 AI ACCESS CONTROL (NEW - SAFE)
# =====================================================

def check_ai_access(user):
    """
    Allow AI access if:
    - User is authenticated
    - User is superuser OR
    - User belongs to an active organization
    - Auto-sync from Profile if membership missing
    """

    if not user.is_authenticated:
        return False, "Login required"

    # ✅ Superusers always allowed
    if user.is_superuser:
        return True, None

    # =========================
    # 🔄 AUTO-SYNC FROM PROFILE
    # =========================
    profile = getattr(user, "profile", None)

    member = (
        OrganizationMember.objects
        .filter(user=user, is_active=True)
        .select_related("organization")
        .first()
    )

    # If no membership exists but profile has organization → auto-create it
    if not member and profile and profile.organization:
        member = OrganizationMember.objects.create(
            user=user,
            organization=profile.organization,
            role=OrganizationMember.ROLE_ADMIN
            if profile.role == profile.ROLE_ORG_ADMIN
            else OrganizationMember.ROLE_MEMBER,
        )

    # =========================
    # 🔐 FINAL VALIDATION
    # =========================
    if not member:
        return False, "AI access restricted"

    if not member.organization or not member.organization.is_active:
        return False, "Organization is inactive"

    return True, None



# =====================================================
# 💬 CHAT VIEW
# =====================================================

@login_required
def chat_view(request, session_id=None):
    user = request.user

    # =========================
    # 🔐 ACCESS CONTROL
    # =========================
    allowed, error = check_ai_access(user)
    if not allowed:
        return HttpResponseForbidden(error)

    # =========================
    # 🧠 ENSURE ONBOARDING CHAT
    # =========================
    create_onboarding_chat(user)

    # =========================
    # 💬 FETCH / CREATE SESSION
    # =========================
    sessions = ChatSession.objects.filter(user=user).order_by("-created_at")

    if session_id:
        session = sessions.filter(id=session_id).first()
    else:
        session = sessions.first()

    if not session:
        session = ChatSession.objects.create(user=user)

    messages = session.messages.order_by("created_at")

    # =========================
    # 📄 ACTIVE DOCUMENT (SAFE SCOPE)
    # =========================
    doc_id = request.GET.get("doc")
    active_document = None

    if doc_id:
        active_document = (
            get_accessible_documents(user)
            .filter(id=doc_id)
            .first()
        )

    # =========================
    # ✉️ HANDLE USER MESSAGE
    # =========================
    if request.method == "POST":
        query = request.POST.get("query", "").strip()

        if query:
            # Save user message
            ChatMessage.objects.create(
                session=session,
                role="user",
                content=query,
            )

            # =========================
            # 🔎 RAG RETRIEVAL
            # =========================
            chunks = retrieve_chunks_for_chat(
                user=user,
                question=query,
                document=active_document,  # may be None
                max_chunks=8,
            )

            # =========================
            # 🤖 RAG ANSWER
            # =========================
            answer_md = rag_answer(
                question=query,
                document=active_document,
            )

            answer_html = markdown.markdown(
                answer_md,
                extensions=["extra", "sane_lists"]
            )

            # =========================
            # 💾 SAVE ASSISTANT MESSAGE
            # =========================
            ChatMessage.objects.create(
                session=session,
                role="assistant",
                content=answer_html,
                sources={
                    "documents": [
                        {
                            "id": active_document.id,
                            "title": active_document.display_name,
                        }
                    ] if active_document else [],
                    "chunks": chunks,
                },
            )

        # =========================
        # 🔁 PRESERVE DOCUMENT CONTEXT
        # =========================
        if active_document:
            return redirect(
                f"{reverse('chat_session', args=[session.id])}?doc={active_document.id}"
            )

        return redirect("chat_session", session_id=session.id)

    # =========================
    # 🖥️ RENDER CHAT UI
    # =========================
    return render(
        request,
        "chat/chat.html",
        {
            "sessions": sessions,
            "active_session": session,
            "messages": messages,
            "active_document": active_document,
        },
    )


# =====================================================
# ❌ DELETE CHAT
# =====================================================

@login_required
def delete_chat(request, session_id):
    session = get_object_or_404(ChatSession, id=session_id, user=request.user)
    session.delete()
    return redirect("chat")


# =====================================================
# ➕ START CHAT WITH CONTEXT
# =====================================================

@login_required
@transaction.atomic
def start_chat_with_context(request):
    if request.method != "POST":
        return redirect("document_list")

    document_ids = request.POST.getlist("documents")
    folder_ids = request.POST.getlist("folders")

    session = ChatSession.objects.create(user=request.user)

    documents = UserDocument.objects.none()

    if document_ids:
        documents = documents | UserDocument.objects.filter(id__in=document_ids)

    if folder_ids:
        documents = documents | UserDocument.objects.filter(folder_id__in=folder_ids)

    for doc in documents.distinct():
        ChatContext.objects.get_or_create(
            session=session,
            document=doc
        )

    return redirect("chat", session_id=session.id)


# =====================================================
# 📄 EXPORT CHAT PDF
# =====================================================

@login_required
def export_chat_pdf(request, session_id):
    session = ChatSession.objects.get(id=session_id, user=request.user)
    messages = session.messages.all()

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="chat_{session.id}.pdf"'

    pdf = canvas.Canvas(response)
    y = 800

    for msg in messages:
        pdf.drawString(40, y, f"{msg.role.upper()}:")
        y -= 20

        for line in msg.content.split("\n"):
            pdf.drawString(60, y, line[:120])
            y -= 15

        y -= 20
        if y < 100:
            pdf.showPage()
            y = 800

    pdf.save()
    return response


# =====================================================
# 📄 EXPORT CHAT DOCX
# =====================================================

@login_required
def export_chat_docx(request, session_id):
    session = ChatSession.objects.get(id=session_id, user=request.user)
    messages = session.messages.all()

    doc = Document()
    doc.add_heading(session.title or "Chat Export", level=1)

    for msg in messages:
        p = doc.add_paragraph()
        p.add_run(msg.role.upper()).bold = True

        for line in msg.content.split("\n"):
            doc.add_paragraph(line, style="Normal")

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    response["Content-Disposition"] = f'attachment; filename="chat_{session.id}.docx"'
    doc.save(response)
    return response


# =====================================================
# 🧹 CLEAR ALL CHATS
# =====================================================

@login_required
def clear_all_chats(request):
    if request.method == "POST":
        with transaction.atomic():
            ChatMessage.objects.filter(session__user=request.user).delete()
            ChatSession.objects.filter(user=request.user).delete()

    return redirect("chat_view")


# =====================================================
# 📄 EXPORT SINGLE ANSWER PDF
# =====================================================

@login_required
def export_answer_pdf(request, message_id):
    message = get_object_or_404(
        ChatMessage,
        id=message_id,
        role="assistant"
    )

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer)
    styles = getSampleStyleSheet()
    story = []

    soup = BeautifulSoup(message.content, "html.parser")

    for element in soup.find_all(["h2", "h3", "p", "li"]):
        story.append(Paragraph(element.get_text(), styles["Normal"]))

    doc.build(story)

    buffer.seek(0)
    response = HttpResponse(buffer, content_type="application/pdf")
    response["Content-Disposition"] = "attachment; filename=answer.pdf"

    return response


# =====================================================
# 📄 EXPORT SINGLE ANSWER DOCX
# =====================================================

@login_required
def export_answer_docx(request, message_id):
    message = get_object_or_404(
        ChatMessage,
        id=message_id,
        role="assistant"
    )

    doc = Document()
    soup = BeautifulSoup(message.content, "html.parser")

    for el in soup.find_all(["h2", "h3", "p", "li"]):
        doc.add_paragraph(el.get_text())

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    response["Content-Disposition"] = "attachment; filename=answer.docx"
    doc.save(response)

    return response


# =====================================================
# 📄 EXPORT SELECTED MESSAGES PDF (RESTORED)
# =====================================================

@login_required
def export_selected_messages_pdf(request, session_id):
    session = get_object_or_404(ChatSession, id=session_id, user=request.user)
    ids = request.POST.getlist("message_ids")
    messages = session.messages.filter(id__in=ids)

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = "attachment; filename=selected_messages.pdf"

    pdf = canvas.Canvas(response)
    y = 800

    for msg in messages:
        pdf.drawString(40, y, f"{msg.role.upper()}:")
        y -= 18

        for line in msg.content.split("\n"):
            pdf.drawString(60, y, line[:120])
            y -= 14

        y -= 20
        if y < 100:
            pdf.showPage()
            y = 800

    pdf.save()
    return response

